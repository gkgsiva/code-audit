"""The summary report: five sections (project, database, integrations, batches, dependency
vulnerabilities) as one HTML page with charts, tables and CSVs, and a PowerPoint built on the
corporate template."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

import pytest
from typer.testing import CliRunner

from legacy_audit.cli import app
from legacy_audit.config import load_config
from legacy_audit.report import pptx_templates
from legacy_audit.report.pptx_templates import resolve_template
from legacy_audit.report.summary import SECTIONS, build_summary_report
from legacy_audit.scan import run_scan
from legacy_audit.security import vulns
from tests.conftest import REPO, SAMPLES
from tests.test_functional import ENRICHMENT, _valid_offline
from tests.test_report import SAMPLE_NAMES


@pytest.fixture(autouse=True)
def _no_security_tools(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(vulns, "find_tool", lambda configured, names: None)


def _scan(sample: str, out: Path) -> Path:
    return run_scan(load_config(SAMPLES / sample / "audit.yaml"), output_dir=out).run_dir


@pytest.mark.parametrize("sample", SAMPLE_NAMES)
def test_summary_page_has_five_sections_and_works_offline(sample: str, tmp_path: Path) -> None:
    run_dir = _scan(sample, tmp_path / "out")
    res = build_summary_report(run_dir)
    assert res.page.is_file() and res.pptx is not None, res.message
    _valid_offline(res.page.parent)
    html = res.page.read_text(encoding="utf-8")
    assert [s[0] for s in SECTIONS] == ["project", "data", "integrations", "batch", "security"]
    for key, *_ in SECTIONS:
        assert f'id="{key}"' in html
    assert 'href="summary.pptx" download' in html
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    for section in data["sections"]:
        for table in section["tables"]:  # every table has its CSV, with the same rows
            path = res.page.parent / "data" / f"{table['key']}.csv"
            with path.open(encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
            assert rows[0] == table["columns"] and len(rows) - 1 == len(table["rows"])


def test_summary_figures_match_the_scan(tmp_path: Path) -> None:
    run_dir = _scan("legacy-layout-app", tmp_path / "out")
    res = build_summary_report(run_dir)
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    by_key = {s["key"]: s for s in data["sections"]}
    audit = json.loads((run_dir / "data" / "audit-summary.json").read_text(encoding="utf-8"))
    project = dict(map(tuple, by_key["project"]["kpis"]))
    assert project["Lines of code"] == audit["totals"]["loc_code"] == 148
    modules = next(t for t in by_key["project"]["tables"] if t["key"] == "modules")
    assert sum(r[-1] for r in modules["rows"]) == 148  # module lines add up to the total
    integ = next(t for t in by_key["integrations"]["tables"] if t["key"] == "integrations")
    assert len(integ["rows"]) == 5
    assert any(r[0] == "login / authentication" and "okta" in r[2] for r in integ["rows"])
    batches = next(t for t in by_key["batch"]["tables"] if t["key"] == "batches")
    assert sorted(r[0] for r in batches["rows"]) == ["billing-batch", "settlement-c"]


@pytest.mark.parametrize("choice", ["blue", "purple"])
def test_summary_pptx_uses_the_corporate_template(tmp_path: Path, choice: str) -> None:
    from pptx import Presentation

    run_dir = _scan("legacy-layout-app", tmp_path / "out")
    res = build_summary_report(run_dir, template=choice)
    assert res.pptx is not None and res.template == choice.title(), res.message
    prs = Presentation(str(res.pptx))
    tpl = Presentation(pptx_templates.template_bytes(resolve_template(choice).path))
    assert [lay.name for lay in prs.slide_layouts] == [lay.name for lay in tpl.slide_layouts]
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[:2] == ["COVER_1", "SUMMARY"] and layouts[-1] == "END"
    assert sum(n.startswith("CHAPTER_") for n in layouts) == len(SECTIONS)
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes)
    assert any(sh.has_table for s in prs.slides for sh in s.shapes)
    text = " ".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
    ).lower()
    assert "lorem" not in text and "integrations at a glance" in text


def test_summary_without_templates_still_writes_the_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(pptx_templates, "template_dirs", lambda: [tmp_path / "none"])
    res = build_summary_report(_scan("c-proc-batch", tmp_path / "out"))
    assert res.page.is_file() and res.pptx is None
    assert 'aria-disabled="true"' in res.page.read_text(encoding="utf-8")


def test_cli_builds_the_summary_report(tmp_path: Path) -> None:
    cfg = SAMPLES / "spring-mvc-batch-app" / "audit.yaml"
    out = tmp_path / "out"
    _scan("spring-mvc-batch-app", out)
    r = CliRunner().invoke(app, ["report", "--type", "summary", "-c", str(cfg), "-o", str(out)])
    assert r.exit_code == 0, r.output
    assert "Summary report" in r.output and "Blue template" in r.output
    assert len(list(out.glob("runs/*/reports/summary/summary.pptx"))) == 1


def test_summary_shows_the_copilot_words_marked_ai(tmp_path: Path) -> None:
    run_dir = _scan("struts-oracle-app", tmp_path / "out")
    shutil.copytree(ENRICHMENT / "struts-oracle-app", run_dir / "enrichment")
    res = build_summary_report(run_dir)
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    sec = {s["key"]: s for s in data["sections"]}["data"]
    assert any(p.startswith("Migration approach (AI): ") for p in sec["intro"])
    tables = {t["key"]: t for t in sec["tables"]}
    cols = tables["tables"]["columns"]
    assert cols[3:] == ["Business data (AI)", "Migration (AI)"]
    orders = next(r for r in tables["tables"]["rows"] if r[0] == "ORDERS")
    assert orders[3] == "Order" and orders[4] == "auto"
    plan = {r[0]: r for r in tables["migration-plan"]["rows"]}
    assert plan["PKG_ORDER_PRICING"][1:3] == ["package", "rewrite"]
    assert "Migration plan per object (AI)" in res.page.read_text(encoding="utf-8")


def test_summary_marks_copilot_batch_words(tmp_path: Path) -> None:
    run_dir = _scan("spring-mvc-batch-app", tmp_path / "out")
    shutil.copytree(ENRICHMENT / "spring-mvc-batch-app", run_dir / "enrichment")
    data = json.loads(
        (build_summary_report(run_dir).page.parent / "summary.json").read_text(encoding="utf-8")
    )
    sec = {s["key"]: s for s in data["sections"]}["batch"]
    rows = {r[0]: r for r in sec["tables"][0]["rows"]}
    assert rows["BIL_NIGHTLY_BOX"][3].endswith("(AI)")
    assert not rows["invoiceJob"][3].endswith("(AI)")  # the scheduler's own description
    assert "Copilot analysis" in sec["source_note"]


def test_summary_without_enrichment_has_no_ai_columns(tmp_path: Path) -> None:
    res = build_summary_report(_scan("struts-oracle-app", tmp_path / "out"))
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    sec = {s["key"]: s for s in data["sections"]}["data"]
    assert [t["key"] for t in sec["tables"]].count("migration-plan") == 0
    assert sec["tables"][0]["columns"] == ["Table", "SQL statements", "CRUD"]


# ---- 5 dependency vulnerabilities ---------------------------------------------------------------
BOTH_DC = REPO / "tests/fixtures/security/billing-dependency-check-report.json"


def _security(run_dir: Path) -> tuple[dict, dict, Path]:  # type: ignore[type-arg]
    res = build_summary_report(run_dir)
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    sec = {s["key"]: s for s in data["sections"]}["security"]
    return sec, {t["key"]: t for t in sec["tables"]}, res.page


def test_vulnerabilities_from_both_scanners_are_merged(sample_copy) -> None:  # type: ignore[no-untyped-def]
    app_dir = sample_copy("spring-mvc-batch-app")  # OSV-Scanner report, plus Dependency-Check's
    shutil.copy(BOTH_DC, app_dir / "audit-inputs" / "security")
    run_dir = run_scan(load_config(app_dir / "audit.yaml")).run_dir
    sec, tables, page = _security(run_dir)
    kpis = dict(map(tuple, sec["kpis"]))
    assert (kpis["Vulnerabilities"], kpis["Critical"], kpis["High"]) == (5, 3, 2)
    assert kpis["Found by both scanners"] == 2 and kpis["Vulnerable libraries"] == 4
    found = {(r[0], r[3]): r[6] for r in tables["vulnerabilities"]["rows"]}
    assert found == {
        ("CVE-2016-1000027", "org.springframework:spring-webmvc"): "Dependency-Check",
        ("CVE-2019-17571", "log4j:log4j"): "both",  # GHSA-2qrg-x229-3v8q in OSV: one row
        ("CVE-2022-23305", "log4j:log4j"): "Dependency-Check",
        ("CVE-2020-26945", "org.mybatis:mybatis"): "both",
        ("CVE-2020-25638", "org.hibernate:hibernate-core"): "OSV-Scanner",
    }
    assert tables["vulnerabilities"]["rows"][0][1] == "critical"  # worst first
    log4j = next(r for r in tables["vulnerabilities"]["rows"] if r[0] == "CVE-2019-17571")
    assert "GHSA-2qrg-x229-3v8q" in log4j[7]
    scanners = {r[0]: r for r in tables["scanners"]["rows"]}
    assert scanners["Dependency-Check"][2:] == [
        "2026-10-15", "audit-inputs/security/billing-dependency-check-report.json", 4, 2,
    ]  # fmt: skip
    assert scanners["OSV-Scanner"][4:] == [3, 1]
    by_scanner = {c["key"]: c for c in sec["charts"]}["by-scanner"]
    assert dict(zip(by_scanner["labels"], by_scanner["values"], strict=True)) == {
        "both": 2, "Dependency-Check": 2, "OSV-Scanner": 1,
    }  # fmt: skip
    assert "both: 2, Dependency-Check only: 2, OSV-Scanner only: 1" in " ".join(sec["intro"])
    assert 'role="alert"' not in page.read_text(encoding="utf-8")


def test_vulnerability_slides_in_the_pptx(sample_copy) -> None:  # type: ignore[no-untyped-def]
    from pptx import Presentation

    app_dir = sample_copy("spring-mvc-batch-app")
    shutil.copy(BOTH_DC, app_dir / "audit-inputs" / "security")
    res = build_summary_report(run_scan(load_config(app_dir / "audit.yaml")).run_dir)
    assert res.pptx is not None, res.message
    prs = Presentation(str(res.pptx))
    titles = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    for title in (
        "Dependency vulnerabilities at a glance",
        "Dependency vulnerabilities: figures",
        "Top vulnerabilities",
        "Vulnerable libraries",
        "Vulnerability checks",
    ):
        assert title in titles, title


def test_no_scanner_results_is_said_plainly(sample_copy) -> None:  # type: ignore[no-untyped-def]
    app_dir = sample_copy("struts-oracle-app")
    shutil.rmtree(app_dir / "audit-inputs" / "security")
    sec, tables, page = _security(run_scan(load_config(app_dir / "audit.yaml")).run_dir)
    assert sec["alert"].startswith("No vulnerability check was done")
    assert "does not mean they are safe" in sec["alert"]
    assert dict(map(tuple, sec["kpis"])) == {"Libraries": 5, "Vulnerability check": "not done"}
    assert sec["charts"] == [] and tables["vulnerabilities"]["rows"] == []
    assert tables["vulnerabilities"]["empty"].startswith("Not checked")
    assert [r[1] for r in tables["scanners"]["rows"]] == ["not used (not installed)"] * 2
    html = page.read_text(encoding="utf-8")
    assert 'role="alert"' in html and "Not checked: no Dependency-Check" in html


def test_one_scanner_names_the_other_as_a_gap(tmp_path: Path) -> None:
    sec, tables, _ = _security(_scan("struts-oracle-app", tmp_path / "out"))
    assert sec["alert"] == "" and "Found by both scanners" not in dict(map(tuple, sec["kpis"]))
    intro = " ".join(sec["intro"])
    assert "Checked with Dependency-Check;" in intro
    assert "Not used: OSV-Scanner (not installed)." in intro
    assert {r[6] for r in tables["vulnerabilities"]["rows"]} == {"Dependency-Check"}


def test_no_libraries_means_nothing_to_check(tmp_path: Path) -> None:
    sec, _, _ = _security(_scan("c-proc-batch", tmp_path / "out"))
    assert sec["intro"][0].startswith("No third-party libraries were found")
    assert sec["alert"] == "" and dict(map(tuple, sec["kpis"]))["Vulnerability check"] == (
        "not needed"
    )


def test_scanners_that_failed_say_why(sample_copy, monkeypatch: pytest.MonkeyPatch) -> None:  # type: ignore[no-untyped-def]
    from legacy_audit.model import VulnerabilitySource

    hint = "offline mode but no OSV database downloaded"
    monkeypatch.setattr(vulns, "run_dependency_check", lambda *a: (
        None, VulnerabilitySource(name=vulns.DC, status="failed", detail="exit code 13"),
    ))  # fmt: skip
    monkeypatch.setattr(vulns, "run_osv_scanner", lambda *a: (
        None, VulnerabilitySource(name=vulns.OSV, status="failed", detail=hint),
    ))  # fmt: skip
    app_dir = sample_copy("struts-oracle-app")
    shutil.rmtree(app_dir / "audit-inputs" / "security")
    sec, tables, _ = _security(run_scan(load_config(app_dir / "audit.yaml")).run_dir)
    intro = " ".join(sec["intro"])
    assert "The scanners were found but did not finish" in intro
    assert "install the tools" not in intro and hint in intro
    assert [r[1] for r in tables["scanners"]["rows"]][1] == f"failed ({hint})"


def test_web_services_in_summary_slides_and_technical_report(tmp_path: Path) -> None:
    from legacy_audit.report.tech import build_tech_report

    run_dir = _scan("web-services-app", tmp_path / "out")
    res = build_summary_report(run_dir)
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    sec = {s["key"]: s for s in data["sections"]}["integrations"]
    kpis = dict(sec["kpis"])
    assert (kpis["Web services offered"], kpis["Web services called"]) == (5, 5)
    web = {t["key"]: t for t in sec["tables"]}["web-services"]
    assert "Suggested migration" not in web["columns"]  # no AI analysis yet
    assert len(web["rows"]) == 10
    assert (res.page.parent / "data" / "web-services.csv").exists()
    deck = json.loads((res.page.parent / "deck.json").read_text(encoding="utf-8"))
    assert {"integrations-web-services", "integrations-web-services-described"} <= {
        s["id"] for s in deck["slides"]
    }
    page = build_tech_report(run_dir).parent / "web-services.html"
    assert "Web service operations" in page.read_text(encoding="utf-8")
    with (page.parent / "data" / "web-services.csv").open(encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    assert rows[0][:3] == ["Offered / called", "REST / SOAP", "Web service"] and len(rows) == 11


def test_web_service_story_from_copilot_is_marked_ai(tmp_path: Path) -> None:
    run_dir = _scan("struts-oracle-app", tmp_path / "out")
    shutil.copytree(ENRICHMENT / "struts-oracle-app", run_dir / "enrichment")
    res = build_summary_report(run_dir)
    data = json.loads((res.page.parent / "summary.json").read_text(encoding="utf-8"))
    sec = {s["key"]: s for s in data["sections"]}["integrations"]
    web = {t["key"]: t for t in sec["tables"]}["web-services"]
    assert web["columns"][-1] == "Suggested migration"
    (row,) = web["rows"]
    assert row[7] == "Central pricing: Gets the confirmed price of an order before it is shown. (AI)"
    assert row[8].endswith("(AI)")
