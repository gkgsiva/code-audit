"""Phase 5: deck.json, the HTML management deck and the editable PPTX."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest
from typer.testing import CliRunner

from legacy_audit.cli import app
from legacy_audit.config import Branding, load_config
from legacy_audit.enrich.validate import validate_run
from legacy_audit.report import pptx as pptx_mod
from legacy_audit.report import pptx_templates
from legacy_audit.report.deck_builder import SECTIONS, build_deck
from legacy_audit.report.mgmt import build_mgmt_report
from legacy_audit.report.pptx import build_pptx, render_diagrams
from legacy_audit.report.pptx_templates import TemplateError, list_templates, resolve_template
from legacy_audit.scan import run_scan
from legacy_audit.security import vulns
from tests.conftest import SAMPLES
from tests.test_report import _check, _summary_value

FIXTURES = Path(__file__).parent / "fixtures" / "enrichment" / "struts-oracle-app"
SAMPLE_NAMES = sorted(p.parent.name for p in SAMPLES.glob("*/expected.yaml"))


@pytest.fixture(autouse=True)
def _no_security_tools(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(vulns, "find_tool", lambda configured, names: None)


@pytest.fixture
def enriched(sample_copy) -> Path:  # type: ignore[no-untyped-def]
    """The Struts sample scanned, with the hand-written enrichment validated into it."""
    app_dir = sample_copy("struts-oracle-app")
    run_dir = run_scan(load_config(app_dir / "audit.yaml")).run_dir
    shutil.copytree(FIXTURES, run_dir / "enrichment")
    assert validate_run(run_dir).ok
    return run_dir


def _blocks(deck: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    return [b for s in deck["slides"] for b in s["blocks"] if b["type"] == kind]


def test_deck_covers_every_section_from_facts_and_enrichment(enriched: Path) -> None:
    deck = build_deck(enriched)
    assert {s["section"] for s in deck["slides"]} == {t for _, t, _ in SECTIONS}
    assert not _blocks(deck, "placeholder")
    ids = {s["id"] for s in deck["slides"]}
    assert {"tech-security", "tech-coverage", "plan-waves", "risk-register"} <= ids
    assert {b["chart"] for b in _blocks(deck, "chart")} == {"bar", "doughnut", "line"}
    summary = json.loads((enriched / "data" / "audit-summary.json").read_text(encoding="utf-8"))
    summary["effort"] = json.loads((enriched / "data" / "effort.json").read_text(encoding="utf-8"))
    for kpis in _blocks(deck, "kpis"):
        for k in kpis["items"]:  # every headline number is the scan's
            assert k["value"] == _summary_value(summary, k["metric"]), k
    assert deck["generated_from"]["missing"] == []


def test_missing_enrichment_shows_placeholders(sample_copy) -> None:  # type: ignore[no-untyped-def]
    app_dir = sample_copy("c-proc-batch")
    run_dir = run_scan(load_config(app_dir / "audit.yaml")).run_dir
    deck = build_deck(run_dir)
    holders = _blocks(deck, "placeholder")
    assert holders and all(re.search(r"run /audit-enrich \w+", b["text"]) for b in holders)
    assert {b["stage"] for b in holders} >= {"overview", "risks", "plan", "narratives"}
    assert len(deck["slides"]) == len(build_deck(run_dir)["slides"])  # same skeleton


def test_stale_enrichment_is_flagged_on_its_slides(enriched: Path) -> None:
    jobs = enriched / "data" / "batch_jobs.json"
    jobs.write_text(jobs.read_text(encoding="utf-8") + " ", encoding="utf-8")  # facts changed
    deck = build_deck(enriched)
    flows = next(s for s in deck["slides"] if s["id"] == "fn-flows")
    assert flows["stale"] == ["flows"]
    assert deck["generated_from"]["enrichment"]["flows"] == "stale"


@pytest.mark.parametrize("sample", SAMPLE_NAMES)
def test_html_deck_is_offline_and_matches_the_summary(sample: str, sample_copy) -> None:  # type: ignore[no-untyped-def]
    app_dir = sample_copy(sample)
    run_dir = run_scan(load_config(app_dir / "audit.yaml")).run_dir
    index = build_mgmt_report(run_dir, Branding(primary_color="red;}body{x"))
    c = _check(index)
    assert c.has_doctype and not c.errors and c.stack == []
    for ref in c.refs:
        assert not re.match(r"(?i)^(https?:)?//", ref)
        assert (index.parent / ref).exists(), ref
    html = index.read_text(encoding="utf-8")
    assert "--brand: #1f4e79" in html  # an unsafe colour falls back to the default
    assert html.count('<section class="slide') == len(build_deck(run_dir)["slides"])
    summary = json.loads((run_dir / "data" / "audit-summary.json").read_text(encoding="utf-8"))
    summary["effort"] = json.loads((run_dir / "data" / "effort.json").read_text(encoding="utf-8"))
    for metric, value in re.findall(r'data-metric="([^"]+)" data-value="([^"]*)"', html):
        assert value == str(_summary_value(summary, metric)), metric
    assert (index.parent / "deck.json").is_file()


def test_render_diagrams_sets_png(enriched: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_run(cmd: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
        Path(cmd[cmd.index("-o") + 1]).write_bytes(b"\x89PNG\r\n\x1a\n")
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(pptx_mod.subprocess, "run", fake_run)
    deck = build_deck(enriched)
    mgmt = enriched / "reports" / "mgmt"
    mgmt.mkdir(parents=True)
    n = render_diagrams(deck, mgmt, "mmdc")
    pngs = [b["png"] for b in _blocks(deck, "diagram")]
    assert n == len(pngs) == 2 and all((mgmt / p).is_file() for p in pngs)


def test_pptx_is_skipped_with_a_message_without_templates(
    enriched: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(pptx_templates, "template_dirs", lambda: [tmp_path / "none"])
    result = build_pptx(enriched)
    assert result.path is None and "no PowerPoint template found" in result.message
    assert (enriched / "reports" / "mgmt" / "deck.json").is_file()  # the deck data is still there


def test_the_three_corporate_templates_are_offered() -> None:
    found = list_templates()
    assert [t.label for t in found] == ["Blue", "Orange", "Purple"]
    assert resolve_template(None).label == "Blue"  # the default
    assert resolve_template("orange").id == "RG_TEMPLATE_PPT_HD_ORANGE"
    assert resolve_template("RG_TEMPLATE_PPT_HD_PURPLE.potx").label == "Purple"
    with pytest.raises(TemplateError, match="available: Blue, Orange, Purple"):
        resolve_template("green")
    assert all(pptx_templates.thumbnail(t.path) for t in found)  # previews for the GUI


def _theme_xml(prs: Any) -> bytes:
    master = prs.slide_masters[0].part
    rel = next(r for r in master.rels.values() if r.reltype.endswith("/theme"))
    return bytes(rel.target_part.blob)


@pytest.mark.parametrize("choice", ["blue", "orange", "purple"])
def test_pptx_is_built_strictly_on_the_chosen_template(enriched: Path, choice: str) -> None:
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    result = build_pptx(enriched, template=choice)
    assert result.path is not None, result.message
    assert result.template == choice.title() and f"({choice.title()} template)" in result.message
    deck = json.loads((enriched / "reports" / "mgmt" / "deck.json").read_text(encoding="utf-8"))
    assert deck["pptx_template"]["label"] == choice.title()
    prs = Presentation(str(result.path))
    tpl = Presentation(pptx_templates.template_bytes(resolve_template(choice).path))
    # the template's own master, theme and layouts, and none of its sample slides
    assert _theme_xml(prs) == _theme_xml(tpl)
    assert [lay.name for lay in prs.slide_layouts] == [lay.name for lay in tpl.slide_layouts]
    assert (prs.slide_width, prs.slide_height) == (tpl.slide_width, tpl.slide_height)
    layouts = [s.slide_layout.name for s in prs.slides]
    sections = [s for s in deck["slides"] if s["type"] == "section"]
    content = [s for s in deck["slides"] if s["type"] == "content"]
    assert layouts[:3] == ["COVER_1", "SUMMARY", "TEXT"]  # cover, contents, overview
    assert layouts[-1] == "END"
    assert len(prs.slides) == len(deck["slides"]) + 3
    assert sum(n.startswith("CHAPTER_") for n in layouts) == len(sections)
    assert sum(n in ("TEXT", "TEXT_DOUBLE") for n in layouts) == len(content) + 1
    text = " ".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
    ).lower()
    assert "lorem" not in text and "chapter title" not in text and "main title" not in text
    titles = [s["title"] for s in deck["slides"]]
    assert all(t.lower() in text for t in titles)
    charts, tables, pictures = [], 0, 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                charts.append(shape.chart.chart_type)
            elif shape.has_table:
                tables += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if shape.has_text_frame:  # every run in the template's heading font or inherited
                for p in shape.text_frame.paragraphs:
                    assert all(r.font.name in (None, "+mj-lt") for r in p.runs)
    expected = [b for b in _blocks(deck, "chart") if b["values"]]
    gauges = sum(len(b["items"]) for b in _blocks(deck, "gauges"))  # one doughnut per gauge
    assert gauges and len(charts) == len(expected) + gauges
    assert {XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.DOUGHNUT} <= set(charts)
    assert XL_CHART_TYPE.LINE_MARKERS in charts
    assert tables == len(_blocks(deck, "table")) + len(_blocks(deck, "heatmap"))
    assert pictures == len([b for b in _blocks(deck, "diagram") if b.get("png")])  # no chart images
    assert prs.slides[0].notes_slide.notes_text_frame.text  # speaker notes carried over


def test_pptx_template_comes_from_audit_yaml(enriched: Path) -> None:
    result = build_pptx(enriched, Branding(pptx_template="purple"))
    assert result.template == "Purple"
    assert build_pptx(enriched, Branding(pptx_template="purple"), template="orange").template == (
        "Orange"
    )  # the reports page / --template wins


def test_report_all_command(enriched: Path) -> None:
    cfg = enriched.parents[2] / "audit.yaml"
    res = CliRunner().invoke(app, ["report", "-c", str(cfg), "--type", "all"])
    assert res.exit_code == 0, res.output
    assert "Technical report" in res.output and "Management deck" in res.output
    assert (enriched / "reports" / "mgmt" / "index.html").is_file()
    assert (enriched / "reports" / "deck.pptx").is_file()
    assert "PowerPoint (Blue template)" in res.output
    res = CliRunner().invoke(
        app, ["report", "-c", str(cfg), "--type", "pptx", "--template", "purple"]
    )
    assert res.exit_code == 0 and "PowerPoint (Purple template)" in res.output, res.output
    res = CliRunner().invoke(app, ["report", "-c", str(cfg), "--type", "pptx", "--template", "red"])
    assert res.exit_code == 1 and "'red' not found" in res.output, res.output


def test_tech_report_analysis_page_shows_enrichment(enriched: Path) -> None:
    from legacy_audit.report.tech import build_tech_report

    page = build_tech_report(enriched).parent / "analysis.html"
    c = _check(page)
    assert not c.errors and c.stack == []
    html = page.read_text(encoding="utf-8")
    assert "SQL injection in order search" in html and "Find and view an order" in html
    assert html.count(">valid<") == 14
    assert "vscode://file/" in html  # claims link to their evidence
