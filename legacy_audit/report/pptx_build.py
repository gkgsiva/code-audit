"""Build the management PPTX from deck.json on a corporate template (python-pptx).

The template's own sample slides are dropped and every slide is added from one of its layouts:
the cover, a contents slide (SUMMARY), a chapter divider per section (CHAPTER_1..4 in turn),
TEXT or TEXT_DOUBLE for content and END. Titles, footers and prose go into the layouts'
placeholders, so they take the template's fonts, sizes and colours. Charts are native PowerPoint
charts (bar, doughnut, line) so they stay editable; tables and heat maps are native tables; gauges
are doughnut charts; flows, capability maps, layers and the roadmap are native shapes; other
diagrams are PNG files rendered from Mermaid when mermaid-cli was available, else a note (the
Mermaid source is in the speaker notes). Every colour drawn is one of the template's theme colours
and every font is its heading font. Every number comes from deck.json, which comes from the scan.
"""

from __future__ import annotations

import math
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import ColorFormat
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationT
from pptx.slide import Slide, SlideLayout
from pptx.text.text import TextFrame
from pptx.util import Emu, Pt

from legacy_audit.report.pptx_templates import Template, TemplateError, template_bytes

EMU_IN = 914400
FONT = "+mj-lt"  # the template's heading font: what its own body placeholders use
GAP = 0.3

COVER, SUMMARY, TEXT, DOUBLE, END = "COVER_1", "SUMMARY", "TEXT", "TEXT_DOUBLE", "END"
CHAPTERS = ("CHAPTER_1", "CHAPTER_2", "CHAPTER_3", "CHAPTER_4")
# placeholder idx per layout, as the corporate templates define them
COVER_TITLE, COVER_SUB = 0, 1
FOOTER, TITLE, BODY, BODY2 = 10, 11, 13, 14
CHAPTER_NO, CHAPTER_TITLE = 16, 17
SUMMARY_SLOTS = ((19, 18), (21, 20), (23, 22), (25, 24), (27, 26), (29, 28), (31, 30), (33, 32))
TEXTY = ("bullets", "text", "placeholder")

T = MSO_THEME_COLOR
SERIES = (
    T.ACCENT_1,
    T.ACCENT_3,
    T.ACCENT_2,
    T.ACCENT_6,
    T.ACCENT_5,
    T.TEXT_2,
    T.ACCENT_4,
    T.BACKGROUND_2,
)
SEMANTIC: dict[str, tuple[MSO_THEME_COLOR, float]] = {
    "accent": (T.ACCENT_1, 0),
    "text": (T.TEXT_1, 0),
    "muted": (T.TEXT_1, 0.4),
    "info": (T.TEXT_1, 0.5),
    "grid": (T.BACKGROUND_1, -0.15),
    "surface": (T.BACKGROUND_1, -0.05),
    "good": (T.ACCENT_2, 0),
    "warning": (T.ACCENT_6, 0),
    "serious": (T.ACCENT_3, 0),
    "critical": (T.ACCENT_5, 0),
}
SEQ = (0.9, 0.75, 0.55, 0.35, 0.0, -0.25, -0.5)  # seq-0..6: light to dark accent 1
RISK = ("good", "warning", "serious", "critical")  # risk-0..3
KIND = {
    "screen": T.ACCENT_1,
    "service": T.ACCENT_4,
    "table": T.ACCENT_2,
    "system": T.ACCENT_3,
    "trigger": T.ACCENT_6,
    "job": T.TEXT_2,
}
FLOW_LEGEND = (
    ("screen", "Screen"),
    ("service", "Service"),
    ("table", "Table"),
    ("system", "Other system"),
    ("trigger", "Trigger"),
    ("job", "Job"),
)


def theme(token: str | None) -> tuple[MSO_THEME_COLOR, float]:
    """A deck.json colour token as one of the template's theme colours (and a tint)."""
    tok = str(token or "accent")
    if tok in SEMANTIC:
        return SEMANTIC[tok]
    head, _, tail = tok.rpartition("-")
    if tail.isdigit():
        n = int(tail)
        if head == "series":
            return SERIES[(n - 1) % len(SERIES)], 0
        if head == "seq":
            return T.ACCENT_1, SEQ[max(0, min(n, len(SEQ) - 1))]
        if head == "risk":
            return SEMANTIC[RISK[max(0, min(n, len(RISK) - 1))]][0], 0.6
    return T.ACCENT_1, 0


def paint(cf: ColorFormat, token: str | None, tint: float | None = None) -> None:
    color, bright = theme(token)
    cf.theme_color = color
    bright = bright if tint is None else tint
    if bright:
        cf.brightness = bright


def kind_color(kind: str) -> MSO_THEME_COLOR:
    return KIND.get(kind, T.TEXT_1)


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float


def E(inches: float) -> Emu:
    return Emu(round(inches * EMU_IN))


def fmt(v: Any) -> str:
    if isinstance(v, bool) or v is None:
        return "" if v is None else str(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.0f}" if v.is_integer() else str(round(v, 1))
    return str(v)


def fit_size(texts: Iterable[str], w: float, h: float, base: int, floor: int = 9) -> int:
    """The largest size (pt) at which the paragraphs roughly fit a w x h inch box."""
    texts = list(texts)
    for size in range(base, floor - 1, -1):
        per_line = max(1, int(w * 72 / (size * 0.52)))
        lines = sum(max(1, math.ceil(len(t) / per_line)) for t in texts)
        if lines * size * 1.25 / 72 + len(texts) * 6 / 72 <= h:
            return size
    return floor


def style_run(
    run: Any,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    tint: float | None = None,
    white: bool = False,
) -> None:
    f = run.font
    f.name = FONT
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if white:
        f.color.theme_color = T.BACKGROUND_1
    elif color:
        paint(f.color, color, tint)


def write(
    tf: TextFrame,
    paras: list[tuple[str, dict[str, Any]]],
    *,
    wrap: bool = True,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float | None = None,
) -> None:
    """Fill a text frame: one paragraph per item, each with its run options."""
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    if margin is not None:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = E(margin)
    for i, (text, opts) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if opts.get("align"):
            p.alignment = opts["align"]
        if opts.get("level") is not None:
            p.level = opts["level"]
        if opts.get("space"):
            p.space_after = Pt(opts["space"])
        run = p.add_run()
        run.text = text
        style_run(
            run,
            opts.get("size"),
            opts.get("bold"),
            opts.get("color"),
            opts.get("tint"),
            opts.get("white", False),
        )


class Builder:
    def __init__(self, deck: dict[str, Any], template: Template, base_dir: Path) -> None:
        self.deck = deck
        self.base_dir = base_dir
        self.prs: PresentationT = Presentation(template_bytes(template.path))
        names = {lay.name: lay for lay in self.prs.slide_layouts}
        needed = (COVER, SUMMARY, TEXT, DOUBLE, END, *CHAPTERS)
        missing = [n for n in needed if n not in names]
        if missing:
            raise TemplateError(
                f"template {template.path.name} has no slide layout(s) named {', '.join(missing)}"
            )
        self.layouts = names
        self.footer = f"{deck['application']} · run {deck['run_id']}"
        self._drop_sample_slides()

    # -- package ------------------------------------------------------------------------------
    def _drop_sample_slides(self) -> None:
        part = self.prs.part
        ids = self.prs.slides._sldIdLst
        for sld in list(ids):
            rid = sld.rId
            ids.remove(sld)
            part.drop_rel(rid)
        for rid, rel in list(part.rels.items()):  # revision history of the sample slides
            if rel.reltype.endswith("/changesInfo"):
                part.drop_rel(rid)

    def _slide(self, layout: str) -> Slide:
        lay: SlideLayout = self.layouts[layout]
        slide = self.prs.slides.add_slide(lay)
        for ph in lay.placeholders:  # python-pptx leaves the footer out; the template has one
            if ph.placeholder_format.idx == FOOTER:
                slide.shapes.clone_placeholder(ph)
        return slide

    @staticmethod
    def _ph(slide: Slide, idx: int) -> Any:
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == idx:
                return ph
        raise TemplateError(f"layout {slide.slide_layout.name} has no placeholder {idx}")

    @staticmethod
    def _drop(shape: Any) -> None:
        el = shape._element
        el.getparent().remove(el)

    def _text_ph(self, slide: Slide, idx: int, text: str) -> None:
        ph = self._ph(slide, idx)
        ph.text_frame.paragraphs[0].add_run().text = text

    def _footer(self, slide: Slide, stale: list[str] | None = None) -> None:
        text = self.footer + (f"   ⚠ Enrichment out of date: {', '.join(stale)}" if stale else "")
        self._text_ph(slide, FOOTER, text)

    @staticmethod
    def _notes(slide: Slide, text: str) -> None:
        if text.strip():
            slide.notes_slide.notes_text_frame.text = text

    def _region(self, layout: str, idx: int = BODY) -> Box:
        for ph in self.layouts[layout].placeholders:
            if ph.placeholder_format.idx == idx:
                return Box(ph.left / EMU_IN, ph.top / EMU_IN, ph.width / EMU_IN, ph.height / EMU_IN)
        raise TemplateError(f"layout {layout} has no placeholder {idx}")

    # -- slides -------------------------------------------------------------------------------
    def build(self, out: Path) -> int:
        slides = self.deck["slides"]
        title = next((s for s in slides if s["type"] == "title"), None)
        self.cover(title)
        self.contents()
        if title and title["blocks"]:
            self.content(dict(title, title="Overview", subtitle=""))
        chapter = {sec["title"]: i for i, sec in enumerate(self.deck.get("sections", []))}
        for s in slides:
            if s["type"] == "title":
                continue
            if s["type"] == "section":
                self.section(s, chapter.get(s["section"], 0))
            else:
                self.content(s)
        self._slide(END)
        props = self.prs.core_properties
        props.title = f"{self.deck['application']}: legacy audit"
        props.author = props.last_modified_by = "legacy-audit"
        props.subject = f"Run {self.deck['run_id']}"
        self.prs.save(str(out))
        return len(self.prs.slides)

    def cover(self, s: dict[str, Any] | None) -> None:
        slide = self._slide(COVER)
        self._text_ph(slide, COVER_TITLE, s["title"] if s else self.deck["application"])
        self._text_ph(slide, COVER_SUB, (s or {}).get("subtitle") or self.footer)
        self._notes(slide, (s or {}).get("notes", ""))

    def contents(self) -> None:
        slide = self._slide(SUMMARY)
        self._text_ph(slide, TITLE, "Contents")
        self._footer(slide)
        sections = self.deck.get("sections", [])
        for i, (no, name) in enumerate(SUMMARY_SLOTS):
            if i >= len(sections):
                self._drop(self._ph(slide, no))
                self._drop(self._ph(slide, name))
                continue
            sec = sections[i]
            self._text_ph(slide, no, f"{i + 1:02d}")
            self._ph(slide, no).text_frame.word_wrap = False  # keep "02" on one line
            tf = self._ph(slide, name).text_frame
            tf.paragraphs[0].add_run().text = sec["title"]
            p = tf.add_paragraph()
            p.level = 1
            p.add_run().text = f"For {sec['audience'].lower()}"
        if len(sections) > len(SUMMARY_SLOTS):
            more = ", ".join(s["title"] for s in sections[len(SUMMARY_SLOTS) :])
            self._notes(slide, f"Also in this deck: {more}.")

    def section(self, s: dict[str, Any], n: int) -> None:
        slide = self._slide(CHAPTERS[n % len(CHAPTERS)])
        self._text_ph(slide, CHAPTER_NO, f"{n + 1:02d}")
        tf = self._ph(slide, CHAPTER_TITLE).text_frame
        tf.paragraphs[0].add_run().text = s["title"]
        if s.get("subtitle"):
            p = tf.add_paragraph()
            p.level = 1
            p.add_run().text = s["subtitle"]
        self._notes(slide, s.get("notes") or "")

    def content(self, s: dict[str, Any]) -> None:
        kpis = [b for b in s["blocks"] if b["type"] == "kpis"]
        rest = [b for b in s["blocks"] if b["type"] != "kpis"]
        texty = [b for b in rest if b["type"] in TEXTY]
        layout = DOUBLE if len(texty) >= 2 else TEXT
        slide = self._slide(layout)
        self._text_ph(slide, TITLE, s["title"])
        self._footer(slide, s.get("stale"))
        region = self._region(TEXT)
        y = region.y
        bottom = region.y + region.h
        if s.get("subtitle"):
            tb = slide.shapes.add_textbox(E(region.x), E(y), E(region.w), E(0.35))
            write(tb.text_frame, [(s["subtitle"], {"size": 14, "color": "muted"})], margin=0)
            y += 0.45
        for b in kpis:
            draw_kpis(slide, b, Box(region.x, y, region.w, 1.0))
            y += 1.4 if b.get("note") else 1.2
        free = [BODY, BODY2] if layout == DOUBLE else [BODY]  # prose goes in these, in order
        if rest:
            n = len(rest)
            share = [1.7, 1.0] if s.get("layout") == "wide-left" and n == 2 else [1.0] * n
            unit = (region.w - GAP * (n - 1)) / sum(share)
            x = region.x
            for b, part in zip(rest, share, strict=True):
                box = Box(x, y, unit * part, bottom - y)
                if b["type"] in TEXTY and free:
                    ph = self._ph(slide, free.pop(0))
                    ph.left, ph.top, ph.width, ph.height = E(box.x), E(box.y), E(box.w), E(box.h)
                    fill_texty(ph.text_frame, b, box)
                else:
                    self.block(slide, b, box)
                x += unit * part + GAP
        for idx in free:
            self._drop(self._ph(slide, idx))
        notes = [s.get("notes") or ""]
        for b in s["blocks"]:
            if b["type"] == "diagram" and b.get("mermaid") and not b.get("png"):
                notes.append(f'Diagram "{b.get("title", "")}" (Mermaid source):\n{b["mermaid"]}')
        if s.get("stale"):
            notes.append(f"Out-of-date enrichment: {', '.join(s['stale'])}.")
        self._notes(slide, "\n\n".join(n for n in notes if n))

    def block(self, slide: Slide, b: dict[str, Any], box: Box) -> None:
        t = b["type"]
        if t in TEXTY:
            tb = slide.shapes.add_textbox(E(box.x), E(box.y), E(box.w), E(box.h))
            fill_texty(tb.text_frame, b, box, textbox=True)
        elif t == "table":
            draw_table(slide, b, box)
        elif t == "chart":
            draw_chart(slide, b, box)
        elif t == "gauges":
            draw_gauges(slide, b, box)
        elif t == "capmap":
            draw_capmap(slide, b, box)
        elif t == "heatmap":
            draw_heatmap(slide, b, box)
        elif t == "flow":
            draw_flow(slide, b, box)
        elif t == "gantt":
            draw_gantt(slide, b, box)
        elif t == "layers":
            draw_layers(slide, b, box)
        elif t == "diagram":
            self.draw_diagram(slide, b, box)

    def draw_diagram(self, slide: Slide, b: dict[str, Any], box: Box) -> None:
        dy = heading(slide, b.get("title"), box)
        png = (self.base_dir / b["png"]) if b.get("png") else None
        if png is not None and png.is_file():
            from PIL import Image

            with Image.open(png) as im:
                pw, ph = im.size
            w, h = box.w, box.h - dy
            scale = min(w / pw, h / ph)
            iw, ih = pw * scale, ph * scale
            slide.shapes.add_picture(str(png), E(box.x + (w - iw) / 2), E(box.y + dy), E(iw), E(ih))
            return
        text = (
            "Diagram not embedded: install @mermaid-js/mermaid-cli (mmdc) and rebuild. The "
            "Mermaid source is in the speaker notes; the HTML deck shows the diagram."
            if b.get("mermaid")
            else (b.get("note") or "No diagram.")
        )
        note_box(slide, Box(box.x, box.y + dy, box.w, 1.2), [(text, {"size": 12})])


# -- blocks -----------------------------------------------------------------------------------
def fill_texty(tf: TextFrame, b: dict[str, Any], box: Box, textbox: bool = False) -> None:
    """Prose in the template's body styles: level 0 heading, 1 paragraph, 2 bullet."""
    t = b["type"]
    paras: list[tuple[str, dict[str, Any]]] = []
    if t == "bullets":
        if b.get("title"):
            paras.append((b["title"], {"level": 0}))
        items = [str(i) for i in b.get("items") or []] or ["None."]
        paras += [(i, {"level": 2}) for i in items]
    elif t == "text":
        paras.append((str(b.get("text", "")), {"level": 1}))
    else:  # placeholder: a stage not enriched yet
        paras += [("Not enriched yet", {"level": 0}), (str(b.get("text", "")), {"level": 1})]
    body = [p for p, o in paras if o["level"]]
    heads = len(paras) - len(body)
    size = fit_size(body, box.w - 0.3, box.h - heads * 0.5, 16 if t == "text" else 14)
    for _, opts in paras:
        if opts["level"]:
            opts["size"] = size
        elif textbox:
            opts.update(size=16, bold=True, color="accent")
        if textbox:
            opts["space"] = 6
            opts.pop("level")
    write(tf, paras, margin=0 if textbox else None)


def heading(slide: Slide, text: str | None, box: Box) -> float:
    if not text:
        return 0
    tb = slide.shapes.add_textbox(E(box.x), E(box.y), E(box.w), E(0.35))
    write(tb.text_frame, [(text.upper(), {"size": 13, "bold": True, "color": "accent"})], margin=0)
    return 0.42


def note_box(slide: Slide, box: Box, paras: list[tuple[str, dict[str, Any]]]) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(box.x), E(box.y), E(box.w), E(box.h))
    shape.fill.solid()
    paint(shape.fill.fore_color, "surface")
    paint(shape.line.color, "grid")
    shape.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    for _, o in paras:
        o.setdefault("color", "text")
    write(shape.text_frame, paras, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)


def rect(
    slide: Slide,
    box: Box,
    fill: tuple[MSO_THEME_COLOR, float],
    line: tuple[MSO_THEME_COLOR, float] | None = None,
    shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
    radius: float | None = None,
) -> Any:
    sh = slide.shapes.add_shape(shape, E(box.x), E(box.y), E(box.w), E(box.h))
    sh.fill.solid()
    sh.fill.fore_color.theme_color = fill[0]
    if fill[1]:
        sh.fill.fore_color.brightness = fill[1]
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.theme_color = line[0]
        if line[1]:
            sh.line.color.brightness = line[1]
        sh.line.width = Pt(1)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def textbox(slide: Slide, box: Box, text: str, **opts: Any) -> Any:
    tb = slide.shapes.add_textbox(E(box.x), E(box.y), E(box.w), E(box.h))
    anchor = opts.pop("anchor", MSO_ANCHOR.TOP)
    write(tb.text_frame, [(text, opts)], anchor=anchor, margin=0)
    return tb


def draw_kpis(slide: Slide, b: dict[str, Any], box: Box) -> None:
    items = b.get("items") or []
    gap = 0.15
    w = (box.w - gap * (len(items) - 1)) / max(len(items), 1)
    for i, k in enumerate(items):
        x = box.x + i * (w + gap)
        rect(slide, Box(x, box.y, w, box.h), SEMANTIC["surface"])
        textbox(
            slide,
            Box(x + 0.12, box.y + 0.08, w - 0.24, 0.5),
            fmt(k["value"]),
            size=24,
            bold=True,
            color="accent",
        )
        label = k["label"] + (f" ({k['unit']})" if k.get("unit") else "")
        textbox(
            slide,
            Box(x + 0.12, box.y + 0.6, w - 0.24, 0.35),
            label,
            size=fit_size([label], w - 0.24, 0.35, 11, 8),
            color="muted",
        )
    if b.get("note"):
        textbox(
            slide, Box(box.x, box.y + box.h + 0.05, box.w, 0.25), b["note"], size=10, color="muted"
        )


def _cell(
    cell: Any,
    text: str,
    size: float,
    bold: bool = False,
    fill: tuple[MSO_THEME_COLOR, float] | None = None,
    white: bool = False,
    align: PP_ALIGN | None = None,
) -> None:
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    style_run(run, size, bold, None if white else "text", white=white)
    cell.margin_left = cell.margin_right = E(0.05)
    cell.margin_top = cell.margin_bottom = E(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.theme_color = fill[0]
        if fill[1]:
            cell.fill.fore_color.brightness = fill[1]


def draw_table(slide: Slide, b: dict[str, Any], box: Box) -> None:
    dy = heading(slide, b.get("title"), box)
    cols = [str(c) for c in b["columns"]]
    rows = [[fmt(c) for c in r] for r in b.get("rows") or []] or [
        ["None."] + [""] * (len(cols) - 1)
    ]
    size = 10 if len(cols) > 4 or len(rows) > 7 else 12
    row_h = size * 1.9 / 72
    weight = [
        math.sqrt(max(4, len(c), *(len(r[i]) if i < len(r) else 0 for r in rows)))
        for i, c in enumerate(cols)
    ]
    total = sum(weight)
    gf = slide.shapes.add_table(
        len(rows) + 1, len(cols), E(box.x), E(box.y + dy), E(box.w), E(row_h * (len(rows) + 1))
    )
    tbl = gf.table  # the built-in style follows the template's accent 1
    for i, wgt in enumerate(weight):
        tbl.columns[i].width = E(box.w * wgt / total)
    for i, c in enumerate(cols):
        _cell(tbl.cell(0, i), c, size, bold=True, white=True)
    for r, row in enumerate(rows, start=1):
        tbl.rows[r].height = E(row_h)
        for i in range(len(cols)):
            _cell(tbl.cell(r, i), row[i] if i < len(row) else "", size)
    tbl.rows[0].height = E(row_h)
    if b.get("more"):
        textbox(
            slide,
            Box(box.x, box.y + box.h - 0.3, box.w, 0.3),
            f"+{b['more']} more in the technical report",
            size=10,
            color="muted",
        )


def _axis_font(axis: Any) -> None:
    tl = axis.tick_labels
    tl.font.size = Pt(10)
    tl.font.name = FONT
    paint(tl.font.color, "muted")


def _hole(plot: Any, size: int) -> None:
    el = plot._element.find(qn("c:holeSize"))
    if el is None:
        el = plot._element.makeelement(qn("c:holeSize"), {})
        plot._element.append(el)
    el.set("val", str(size))


def draw_chart(slide: Slide, b: dict[str, Any], box: Box) -> None:
    dy = heading(slide, b.get("title"), box)
    values = b.get("values") or []
    if not values:
        textbox(
            slide, Box(box.x, box.y + dy, box.w, 0.4), "Nothing to chart.", size=12, color="muted"
        )
        return
    data = CategoryChartData()
    data.categories = [str(v) for v in b["labels"]]
    data.add_series(b.get("title") or "Values", values)
    kind = b.get("chart")
    ctype = {
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "line": XL_CHART_TYPE.LINE_MARKERS,
    }.get(kind or "", XL_CHART_TYPE.COLUMN_CLUSTERED)
    h = min(box.h - dy, 4.6)
    chart = slide.shapes.add_chart(ctype, E(box.x), E(box.y + dy), E(box.w), E(h), data).chart
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(10)
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(10)
    labels.font.name = FONT
    labels.number_format_is_linked = True
    series = plot.series[0]
    colors = b.get("colors") or []
    if kind == "doughnut":
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        labels.show_value = True
        labels.font.color.theme_color = T.BACKGROUND_1
        _hole(plot, 55)
        for i, point in enumerate(series.points):
            point.format.fill.solid()
            paint(point.format.fill.fore_color, colors[i] if i < len(colors) else f"series-{i + 1}")
        return
    chart.has_legend = False
    _axis_font(chart.category_axis)
    _axis_font(chart.value_axis)
    chart.category_axis.has_major_gridlines = False
    grid = chart.value_axis.major_gridlines.format.line
    paint(grid.color, "grid")
    grid.width = Pt(0.5)
    chart.value_axis.format.line.fill.background()
    if b.get("unit"):
        chart.value_axis.has_title = True
        tf = chart.value_axis.axis_title.text_frame
        tf.text = str(b["unit"])
        style_run(tf.paragraphs[0].runs[0], 10, False, "muted")
    paint(labels.font.color, "text")
    if kind == "line":
        labels.position = XL_LABEL_POSITION.ABOVE
        series.format.line.width = Pt(2)
        paint(series.format.line.color, "accent")
        series.smooth = False
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 7
        series.marker.format.fill.solid()
        paint(series.marker.format.fill.fore_color, "accent")
        return
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.gap_width = 60
    plot.vary_by_categories = False
    if colors and all(c == colors[0] for c in colors):
        series.format.fill.solid()
        paint(series.format.fill.fore_color, colors[0])
    else:
        for i, point in enumerate(series.points):
            point.format.fill.solid()
            paint(point.format.fill.fore_color, colors[i] if i < len(colors) else "accent")


def draw_gauges(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """Share of a whole as a ring: a native doughnut chart with the percentage in the middle."""
    items = b.get("items") or []
    gap = 0.25
    w = min(3.0, (box.w - gap * (len(items) - 1)) / max(len(items), 1))
    left = box.x + (box.w - (w * len(items) + gap * (len(items) - 1))) / 2
    ring = min(w - 0.4, 2.2)
    for i, g in enumerate(items):
        x = left + i * (w + gap)
        rect(
            slide,
            Box(x, box.y, w, ring + 1.3),
            SEMANTIC["surface"],
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.05,
        )
        pct = max(0, min(100, g["pct"]))
        data = CategoryChartData()
        data.categories = ["Share", "Rest"]
        data.add_series(g["label"], [pct, 100 - pct])
        rx = x + (w - ring) / 2
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, E(rx), E(box.y + 0.15), E(ring), E(ring), data
        ).chart
        chart.has_title = False
        chart.has_legend = False
        plot = chart.plots[0]
        plot.has_data_labels = False
        _hole(plot, 72)
        for k, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            paint(point.format.fill.fore_color, g.get("color") if k == 0 else "grid")
            point.format.line.fill.background()
        textbox(
            slide,
            Box(rx, box.y + 0.15 + ring / 2 - 0.3, ring, 0.6),
            f"{pct}%",
            size=22,
            bold=True,
            color="text",
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        textbox(
            slide,
            Box(x + 0.1, box.y + ring + 0.25, w - 0.2, 0.55),
            g["label"],
            size=fit_size([g["label"]], w - 0.2, 0.55, 12),
            bold=True,
            color="text",
            align=PP_ALIGN.CENTER,
        )
        if g.get("sub"):
            textbox(
                slide,
                Box(x + 0.1, box.y + ring + 0.8, w - 0.2, 0.35),
                g["sub"],
                size=fit_size([g["sub"]], w - 0.2, 0.35, 10, 8),
                color="muted",
                align=PP_ALIGN.CENTER,
            )


def _shape_text(
    sh: Any,
    text: str,
    size: float,
    *,
    bold: bool = False,
    white: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    margin: float = 0.05,
) -> None:
    write(
        sh.text_frame,
        [(text, {"size": size, "bold": bold, "white": white, "color": "text", "align": align})],
        anchor=MSO_ANCHOR.MIDDLE,
        margin=margin,
    )


def draw_capmap(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """Domains as tinted groups; each capability a white tile with the domain's colour edge."""
    domains = b.get("domains") or []
    cols = min(max(len(domains), 1), 4)
    gap = 0.2
    w = (box.w - gap * (cols - 1)) / cols
    nrows = math.ceil(len(domains) / cols) or 1
    row_h = (box.h - gap * (nrows - 1)) / nrows
    for i, d in enumerate(domains):
        color = theme(d.get("color"))[0]
        x = box.x + (i % cols) * (w + gap)
        y = box.y + (i // cols) * (row_h + gap)
        caps = d.get("caps") or []
        cap_h = min(0.42, (row_h - 0.6) / max(len(caps), 1))
        h = min(row_h, 0.6 + len(caps) * (cap_h + 0.06) + 0.1)
        rect(
            slide,
            Box(x, y, w, h),
            (color, 0.85),
            (color, 0.5),
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.04,
        )
        textbox(
            slide,
            Box(x + 0.15, y + 0.08, w - 0.3, 0.42),
            d["name"],
            size=fit_size([d["name"]], w - 0.3, 0.42, 14, 10),
            bold=True,
            color="text",
        )
        for k, cap in enumerate(caps):
            cy = y + 0.55 + k * (cap_h + 0.06)
            tile = rect(slide, Box(x + 0.15, cy, w - 0.3, cap_h), (T.BACKGROUND_1, 0))
            _shape_text(tile, cap, fit_size([cap], w - 0.5, cap_h, 11, 8), margin=0.12)
            rect(slide, Box(x + 0.15, cy, 0.06, cap_h), (color, 0))


def draw_heatmap(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """A native table whose cell fill shows how much (seq) or how risky (risk); the text stays."""
    dy = heading(slide, b.get("title"), box)
    scale = "risk" if b.get("scale") == "risk" else "seq"
    cols = [str(c) for c in b["columns"]]
    rows = b.get("rows") or []
    n = max(len(rows), 1)
    first = min(3.0, box.w * 0.28)
    rest = (box.w - first) / max(len(cols), 1)
    gf = slide.shapes.add_table(
        n + 1, len(cols) + 1, E(box.x), E(box.y + dy), E(box.w), E(0.34 * (n + 1))
    )
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl.columns[0].width = E(first)
    for i in range(len(cols)):
        tbl.columns[i + 1].width = E(rest)
    white = (T.BACKGROUND_1, 0.0)
    _cell(tbl.cell(0, 0), "", 10, fill=white)
    for i, c in enumerate(cols):
        _cell(tbl.cell(0, i + 1), c, 10, bold=True, fill=white, align=PP_ALIGN.CENTER)
    if not rows:
        _cell(tbl.cell(1, 0), "None.", 10, fill=white)
        for i in range(len(cols)):
            _cell(tbl.cell(1, i + 1), "", 10, fill=white)
    for r, row in enumerate(rows, start=1):
        _cell(tbl.cell(r, 0), str(row["label"]), 10, bold=True, fill=white)
        for i, c in enumerate(row["cells"]):
            level = int(c.get("level", 0))
            fill = theme(f"{scale}-{level}")
            dark = scale == "seq" and level >= 4
            _cell(
                tbl.cell(r, i + 1),
                str(c.get("text") or ""),
                10,
                bold=True,
                fill=fill,
                white=dark,
                align=PP_ALIGN.CENTER,
            )
    for r in range(n + 1):
        tbl.rows[r].height = E(0.34)
    if b.get("legend"):
        ty = min(box.y + dy + 0.34 * (n + 1) + 0.1, box.y + box.h - 0.3)
        textbox(slide, Box(box.x, ty, box.w, 0.3), b["legend"], size=10, color="muted")


def _arrow(slide: Slide, x1: float, y1: float, x2: float, y2: float) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    paint(c.line.color, "muted")
    c.line.width = Pt(1)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))


def draw_flow(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """A flow as columns of coloured boxes joined by arrows (all editable shapes)."""
    cols = [c for c in b.get("columns") or [] if c.get("nodes")]
    top, legend_h = box.y + 0.35, 0.3
    h = box.h - 0.35 - legend_h
    col_w = box.w / max(len(cols), 1)
    node_w = min(2.6, col_w * 0.74)
    pos: dict[str, tuple[int, Box]] = {}
    nodes: list[tuple[str, str, Box]] = []
    for i, c in enumerate(cols):
        cx = box.x + i * col_w + (col_w - node_w) / 2
        textbox(
            slide,
            Box(box.x + i * col_w, box.y, col_w, 0.3),
            c["title"],
            size=10,
            bold=True,
            color="muted",
            align=PP_ALIGN.CENTER,
        )
        slot = h / len(c["nodes"])
        node_h = min(0.62, slot * 0.72)
        for k, nd in enumerate(c["nodes"]):
            nb = Box(cx, top + k * slot + (slot - node_h) / 2, node_w, node_h)
            nodes.append((str(nd["label"]), str(nd.get("kind", "")), nb))
            pos.setdefault(nd["key"], (i, nb))
    labels: list[tuple[str, Box]] = []
    for e in b.get("edges") or []:  # arrows first, then the boxes on top of them, then labels
        if e["from"] not in pos or e["to"] not in pos or e["from"] == e["to"]:
            continue
        (ca, a), (cz, z) = pos[e["from"]], pos[e["to"]]
        if ca == cz:
            x1, x2 = a.x + a.w * 0.85, z.x + z.w * 0.85
            y1, y2 = (a.y + a.h, z.y) if z.y > a.y else (a.y, z.y + z.h)
        elif cz > ca:
            x1, y1, x2, y2 = a.x + a.w, a.y + a.h / 2, z.x, z.y + z.h / 2
        else:
            x1, y1, x2, y2 = a.x, a.y + a.h / 2, z.x + z.w, z.y + z.h / 2
        _arrow(slide, x1, y1, x2, y2)
        if e.get("label"):
            lw = min(1.6, 0.12 + len(str(e["label"])) * 0.065)
            mx, my = x1 + (x2 - x1) * 0.55, y1 + (y2 - y1) * 0.55
            labels.append((str(e["label"]), Box(mx - lw / 2, my - 0.11, lw, 0.22)))
    for label, kind, nb in nodes:
        color = kind_color(kind)
        shape = (
            MSO_SHAPE.CAN
            if kind == "table"
            else MSO_SHAPE.ROUNDED_RECTANGLE
            if kind in ("system", "trigger")
            else MSO_SHAPE.RECTANGLE
        )
        sh = rect(slide, nb, (color, 0.8), (color, 0), shape=shape)
        size = 8 if nb.h < 0.4 else 10
        _shape_text(sh, label, fit_size([label], nb.w - 0.1, nb.h, size, 7), align=PP_ALIGN.CENTER)
    for label, lb in labels:
        sh = rect(slide, lb, (T.BACKGROUND_1, 0))
        _shape_text(sh, label, 8, align=PP_ALIGN.CENTER, margin=0)
    lx = box.x
    for kind, label in FLOW_LEGEND:
        color = kind_color(kind)
        rect(slide, Box(lx, box.y + box.h - 0.22, 0.16, 0.16), (color, 0.8), (color, 0))
        textbox(
            slide, Box(lx + 0.2, box.y + box.h - 0.27, 1.2, 0.26), label, size=10, color="muted"
        )
        lx += 0.35 + len(label) * 0.08


def draw_gantt(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """Plan waves as bars on a month axis."""
    bars = b.get("bars") or []
    span = max(b.get("months") or 0, 1)
    label_w = min(3.0, box.w * 0.25)
    tx, tw = box.x + label_w + 0.15, box.w - label_w - 0.15
    row_h = min(0.6, (box.h - 0.5) / max(len(bars), 1))
    for i, g in enumerate(bars):
        y = box.y + i * row_h
        textbox(
            slide,
            Box(box.x, y, label_w, row_h * 0.85),
            g["name"],
            size=fit_size([g["name"]], label_w, row_h * 0.85, 12, 9),
            bold=True,
            color="text",
            anchor=MSO_ANCHOR.MIDDLE,
        )
        rect(
            slide,
            Box(tx, y + row_h * 0.1, tw, row_h * 0.65),
            SEMANTIC["surface"],
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.2,
        )
        bw = max(0.05, tw * (g["end"] - g["start"]) / span)
        color = theme(g.get("color"))[0]
        bar = rect(
            slide,
            Box(tx + tw * g["start"] / span, y + row_h * 0.1, bw, row_h * 0.65),
            (color, 0),
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.2,
        )
        _shape_text(bar, f"M{g['start']} to M{g['end']}", 10, white=True, margin=0.06)
    step = 1 if span <= 12 else 3 if span <= 36 else 6
    ay = box.y + len(bars) * row_h + 0.05
    for m in range(0, span + 1, step):
        x = tx + tw * m / span
        textbox(
            slide,
            Box(x - 0.3, ay, 0.6, 0.25),
            f"M{m}",
            size=10,
            color="muted",
            align=PP_ALIGN.CENTER,
        )


def draw_layers(slide: Slide, b: dict[str, Any], box: Box) -> None:
    """Architecture layers as tinted bands: name on the left, components and note on the right."""
    layers = b.get("layers") or []
    gap = 0.12
    h = min(1.1, (box.h - gap * (len(layers) - 1)) / max(len(layers), 1))
    name_w = min(2.6, box.w * 0.24)
    for i, layer in enumerate(layers):
        y = box.y + i * (h + gap)
        color = theme(layer.get("color"))[0]
        rect(slide, Box(box.x, y, box.w, h), (color, 0.85), (color, 0.5))
        band = rect(slide, Box(box.x, y, name_w, h), (color, 0))
        _shape_text(
            band,
            layer["name"],
            fit_size([layer["name"]], name_w - 0.2, h, 14, 9),
            bold=True,
            white=True,
            margin=0.1,
        )
        items = "   ·   ".join(layer.get("items") or []) or "(no components)"
        paras: list[tuple[str, dict[str, Any]]] = [(items, {"bold": True, "color": "text"})]
        if layer.get("note"):
            paras.append((layer["note"], {"color": "muted"}))
        size = fit_size([p for p, _ in paras], box.w - name_w - 0.3, h, 12, 8)
        for _, o in paras:
            o["size"] = size
        tb = slide.shapes.add_textbox(
            E(box.x + name_w + 0.15), E(y), E(box.w - name_w - 0.25), E(h)
        )
        write(tb.text_frame, paras, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def build_on_template(deck: dict[str, Any], template: Template, base_dir: Path, out: Path) -> int:
    """Write ``out`` from deck.json on ``template``; return the number of slides."""
    return Builder(deck, template, base_dir).build(out)
